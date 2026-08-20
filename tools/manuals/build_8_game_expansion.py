from pathlib import Path
from copy import deepcopy
import math

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


PROJECT_ROOT = Path(__file__).resolve().parents[2]
OUT = PROJECT_ROOT / "manuals" / "expansion-pack"
COMPILED = OUT / "compiled"
ASSETS = PROJECT_ROOT / "assets" / "expansion-pack"
QA = PROJECT_ROOT / "artifacts" / "qa" / "expansion-pack"
for p in (OUT, COMPILED, ASSETS, QA):
    p.mkdir(parents=True, exist_ok=True)

KIT_IMAGE = PROJECT_ROOT / "assets" / "handbook" / "kit-concept.png"
PLAY_IMAGE = PROJECT_ROOT / "assets" / "handbook" / "gameplay-villa.png"
DOCX = COMPILED / "Eight-New-Games-Expansion-Pack.docx"
PDF = COMPILED / "Eight-New-Games-Expansion-Pack.pdf"

NAVY = "102A43"
TEAL = "0C7C86"
CORAL = "E76F51"
AMBER = "E9A23B"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
GRAY = "5C6770"
PALE = "F4F7FA"
TABLE_FILL = "E8EEF5"


GAMES = [
    {
        "number": 1,
        "name": "Crown Circuit",
        "tagline": "Hold the crown, survive challenges, and time every duel.",
        "tags": ["COMPETITIVE", "TIMING", "MEDIUM MOVEMENT"],
        "meta": [
            ("Players", "4-6"), ("Duration", "8-12 minutes"),
            ("Device position", "Upper arm or front torso"), ("Activity", "Medium"),
            ("Mode", "Free-for-all"), ("Best space", "Clear room or garden")
        ],
        "objective": "Finish with the most Prestige by controlling the digital Crown, stealing it through clean device presses, and winning short reaction duels.",
        "setup": [
            "Assign one pod to each player and attach it where the player can safely press their own pod during a duel.",
            "Clear the play boundary and forbid stairs, roads, pools, door blocking, pushing, grabbing, and contact anywhere except the device.",
            "Choose an 8-, 10-, or 12-minute match and a target of 15 or 20 Prestige.",
            "Run the reaction calibration: every player waits for their LED cue and presses once.",
            "The hub randomly awards the first Crown; that pod pulses amber while all others remain teal."
        ],
        "device": [
            "Amber pulse means Crown holder; teal means challenger; blue means duel pending; red means locked out.",
            "A challenger starts a duel by pressing the Crown holder's physical pod once.",
            "The two duelists step apart. After a random 1-3 second delay, both pods flash; the first valid self-press wins.",
            "A press before the cue is a false start. The hub timestamps both events and decides the result; nodes never decide locally.",
            "All other pods lock during the duel so outside presses cannot affect the result."
        ],
        "rules": [
            "The Crown holder earns one Prestige for every five full seconds of uninterrupted control.",
            "A successful challenger gains two Prestige and receives the Crown.",
            "A successful defense gives the Crown holder two Prestige.",
            "After any duel, the winner receives a 10-second shield shown by a white LED ring; shielded players cannot be challenged.",
            "A player may not challenge the same Crown holder twice in a row unless every other active player has attempted a challenge.",
            "The round ends when the timer expires or a player reaches the target. Highest Prestige wins; longest total Crown time breaks a tie."
        ],
        "strategy": "The obvious move is to chase every opportunity, but timing matters more. Challengers should watch the shield and approach from a safe angle; the Crown holder can force impatient rivals into false starts. The rotation rule prevents two players from farming each other.",
        "variations": [
            "Team Circuit: two teams share Prestige; a Crown transfer within the team requires a two-second hold.",
            "Quiet Crown: use LED and vibration only for indoor gatherings.",
            "No-chase version: pods are placed on a table and players challenge only when the hub calls their name."
        ],
        "safety": "Use walking or controlled movement in small rooms. Only the device may be touched. Pause immediately if the attachment loosens or players begin blocking, grabbing, or sprinting through unsafe areas."
    },
    {
        "number": 2,
        "name": "Courier Code",
        "tagline": "Pass a hidden digital packet while an Interceptor searches the group.",
        "tags": ["SOCIAL DEDUCTION", "HIDDEN ROLES", "LOW MOVEMENT"],
        "meta": [
            ("Players", "5-6"), ("Duration", "10-15 minutes"),
            ("Device position", "Handheld or front clip"), ("Activity", "Low"),
            ("Mode", "Courier team vs Interceptor"), ("Best space", "Majlis, lounge, or table")
        ],
        "objective": "The Courier team must complete three secret packet handoffs and extract before time expires. The hidden Interceptor must identify and scan the current carrier.",
        "setup": [
            "Give every player a pod, then ask players to close their eyes or look away from the host screen.",
            "The hub privately assigns one starting Courier with a long vibration, one Interceptor with two short vibrations, and all other players as Agents with one short vibration.",
            "Set a 12-minute round, three required handoffs, and three Interceptor scan tokens.",
            "Enable private feedback: gameplay role and packet cues use vibration; LEDs remain neutral until a scan or final result."
        ],
        "device": [
            "The current carrier feels a soft heartbeat vibration every 20 seconds; nobody else receives it.",
            "To attempt a handoff, the carrier holds the recipient's pod for 1.5 seconds. The recipient must press their own pod within 2.5 seconds.",
            "A valid handoff transfers the packet and privately vibrates both pods. A decoy attempt produces no confirmation, allowing Agents to bluff.",
            "The Interceptor scans by double-pressing a target's pod within one second. The hub spends one scan token and reports hit or miss.",
            "The host screen shows time, number of completed handoffs, and scans remaining, but never shows the carrier."
        ],
        "rules": [
            "Players may talk, accuse, bluff, and perform decoy handoff gestures, but may not hide or cover another player's pod.",
            "A carrier cannot hand the packet back to the player who gave it to them.",
            "After each real handoff, the new carrier has a 12-second scan shield so the transfer cannot be punished instantly.",
            "After three real handoffs, the carrier may extract by holding their own pod for three seconds during the final two minutes.",
            "A correct scan ends the game for the Interceptor. An incorrect scan consumes one token and publicly marks the scanned player clear for 30 seconds.",
            "If all scan tokens are used, the Interceptor may still win by correctly naming the carrier in one final host-screen accusation."
        ],
        "strategy": "Real and fake handoffs should look alike. The Courier team benefits from several believable decoys, while the Interceptor should track who avoids risk, who suddenly changes behavior, and which handoffs happen near shield timing.",
        "variations": [
            "Open Interceptor: reveal the Interceptor for a lighter family version.",
            "Two packets: for six experienced players, run two carriers and require both to extract.",
            "Silent round: discussion is allowed only during two 45-second briefing windows."
        ],
        "safety": "This is a low-movement game. Handoffs are gentle presses on the device only; no searching clothing, restraining hands, or forcing a player to reveal private feedback."
    },
    {
        "number": 3,
        "name": "Double Lock",
        "tagline": "Solve clues and synchronize the correct pods before the lock resets.",
        "tags": ["COOPERATIVE", "SYNCHRONIZATION", "PUZZLE"],
        "meta": [
            ("Players", "3-6"), ("Duration", "8-12 minutes"),
            ("Device position", "Table or safe stations"), ("Activity", "Low to medium"),
            ("Mode", "Cooperative"), ("Best space", "Table, room, or garden stations")
        ],
        "objective": "Open six digital locks by interpreting clues and holding the correct pair or trio of pods at the same time before the countdown reaches zero.",
        "setup": [
            "Place all six pods where they are easy to reach and label their colors on the host screen.",
            "Choose Table mode for low movement or Station mode for a larger safe area.",
            "Select six locks, a 10-minute master timer, and either two or three allowed mistakes per lock.",
            "Run a synchronization test: two named players hold their pods together for two seconds."
        ],
        "device": [
            "The hub presents a clue such as 'the two colors that make green,' 'the neighbors of amber,' or 'the pair that flashed before blue.'",
            "Candidate pods glow dimly. Players may inspect and discuss without pressing.",
            "A lock attempt begins when one candidate is held; the matching pod must be held within 400 ms.",
            "Both holds must remain active for two seconds. LEDs fill toward the center to show synchronization progress.",
            "Correct combinations turn green and vibrate once. Incorrect combinations flash red, add a strike, and reset after three seconds."
        ],
        "rules": [
            "The team receives one clue at a time and may make only one active attempt at once.",
            "Players may not move a pod after the first lock begins unless the host pauses the game.",
            "Locks 1-2 use direct clues; locks 3-4 use memory; lock 5 adds a false candidate; lock 6 requires three synchronized pods.",
            "Three strikes on one lock trigger a 20-second penalty and reveal a stronger clue.",
            "The team wins by opening all six locks before the master timer expires."
        ],
        "strategy": "Assign one player to read, one to track earlier flashes, and one to count the synchronized hold. In Station mode, agree who covers each area before the clue appears so players do not collide.",
        "variations": [
            "Family clues: colors, simple sums, and visible sequences.",
            "Expert clues: negative clues, remembered order, and timed three-pod holds.",
            "Two-team race: each team controls three pods and must coordinate cross-team pairs for neutral locks."
        ],
        "safety": "Station mode must use a clear one-way movement plan and non-slip surfaces. Place pods away from furniture edges, stairs, water, food preparation, and traffic."
    },
    {
        "number": 4,
        "name": "Market Rush",
        "tagline": "Negotiate digital resources, confirm trades, and complete changing contracts.",
        "tags": ["NEGOTIATION", "RESOURCE STRATEGY", "SEATED"],
        "meta": [
            ("Players", "4-6"), ("Duration", "15-20 minutes"),
            ("Device position", "Handheld or table"), ("Activity", "Low"),
            ("Mode", "Competitive negotiation"), ("Best space", "Table or majlis")
        ],
        "objective": "Earn the most Reputation by trading three fictional resources - Spark, Fabric, and Time - and completing public contracts before the market changes.",
        "setup": [
            "Assign one pod per player and place the shared market dashboard where everyone can see it.",
            "The hub gives each player five random resource units and displays three public contracts.",
            "Choose four 90-second trading rounds plus 20-second settlement windows.",
            "Run one practice trade so everyone understands select, confirm, and cancel."
        ],
        "device": [
            "Tap another player's pod once to open a trade between those two players on the host screen.",
            "Each player short-presses their own pod to cycle the resource they will give; double-press changes quantity between one and two.",
            "A long hold locks that side of the offer. When both players are locked, both pods pulse white.",
            "Both players press once within three seconds to execute. Either player may hold for two seconds to cancel before execution.",
            "A player completes an affordable public contract by holding their own pod during the settlement window; the hub spends resources automatically."
        ],
        "rules": [
            "Players may promise future trades, but only hub-confirmed exchanges are binding.",
            "Only one trade per player may be open at a time; other invitations queue for five seconds and then expire.",
            "Each contract awards Reputation and is replaced immediately after completion.",
            "At the end of every trading round, a market event changes one resource bonus or contract type.",
            "Unspent resources are worth little at the end, so hoarding is rarely optimal.",
            "After four rounds, highest Reputation wins; completed contracts break ties."
        ],
        "strategy": "Information is public but urgency is not. Players can combine two modest trades into a high-value contract or create temporary alliances. The hub prevents accidental or impossible trades, so discussion remains the main skill.",
        "variations": [
            "Cooperative market: the group must complete ten contracts before four rounds end.",
            "Private holdings: show exact inventories only while a player presses and holds their pod near the screen.",
            "Fast market: 60-second rounds and single-unit trades only."
        ],
        "safety": "All resources and Reputation are fictional game values. Do not connect the mechanic to money, betting, prizes based on chance, or real-world financial promises."
    },
    {
        "number": 5,
        "name": "Power Grid",
        "tagline": "Balance output, heat, and emergencies across six connected generators.",
        "tags": ["COOPERATIVE", "REAL-TIME STRATEGY", "TABLE"],
        "meta": [
            ("Players", "3-6"), ("Duration", "8-12 minutes"),
            ("Device position", "Table or floor stations"), ("Activity", "Low to medium"),
            ("Mode", "Cooperative"), ("Best space", "Table or six safe stations")
        ],
        "objective": "Keep total power within the city's changing demand band while preventing any generator from overheating until the shift ends.",
        "setup": [
            "Place all six pods in a visible ring and assign one or two generators to each player.",
            "Select Easy, Standard, or Storm difficulty and an 8- or 10-minute shift.",
            "The host screen displays city demand, current supply, stability, and upcoming forecast.",
            "Practice raising output, cooling, and performing a linked repair."
        ],
        "device": [
            "LED brightness shows generator output from zero to three; color shifts from teal to amber to red as heat increases.",
            "A short press raises output one level. A long hold cools the generator but lowers output one level.",
            "A double press enables a five-second boost: extra output now, followed by extra heat.",
            "When a repair symbol appears, two specified pods must be held simultaneously for two seconds.",
            "An overheated generator locks for 10 seconds and produces no power."
        ],
        "rules": [
            "Demand changes every 12-25 seconds and is previewed by a short forecast unless a Storm event hides it.",
            "Supply more than one unit above or below demand drains Stability; an exact match slowly restores it.",
            "Generators gain heat while producing at level three and cool while at zero or during a long hold.",
            "Crisis events include a maintenance pair, temporary generator outage, rapid demand spike, and controlled shutdown.",
            "The team wins if time expires with Stability above zero. The team loses if Stability reaches zero or three generators overheat together."
        ],
        "strategy": "Use the forecast rather than chasing the current number. Keep one generator cool as reserve, announce every boost, and assign a player to watch Stability instead of a physical pod.",
        "variations": [
            "Young/family mode: slower demand changes and no hidden forecast.",
            "Blackout challenge: begin at zero and reach a stable target through a controlled sequence.",
            "Two-grid mode: teams manage three generators each but share one city Stability meter."
        ],
        "safety": "In floor-station mode, players walk between fixed stations; do not run. Keep pods visible, stable, dry, and outside paths where they could be stepped on."
    },
    {
        "number": 6,
        "name": "Rhythm Rivals",
        "tagline": "Turn every pod into a personal beat controller and synchronize the team.",
        "tags": ["RHYTHM", "TEAM COMPETITION", "LOW MOVEMENT"],
        "meta": [
            ("Players", "2-6"), ("Duration", "6-10 minutes"),
            ("Device position", "Handheld or table"), ("Activity", "Low"),
            ("Mode", "Solo or two teams"), ("Best space", "Seated or standing circle")
        ],
        "objective": "Score accurate notes, complete team sequences, and finish with the highest Groove score using original click-and-percussion tracks from the hub.",
        "setup": [
            "Place each pod in a player's hand or directly in front of them on a stable surface.",
            "Run latency calibration: the hub flashes and clicks five times while players press on cue.",
            "Choose Solo, Teams, or Cooperative; select tempo and a three-, five-, or seven-track set.",
            "Set sound volume appropriate to the room or enable visual-and-vibration mode."
        ],
        "device": [
            "A pod glows one beat before its note. The player presses on the audio click or full-bright LED cue.",
            "Short presses play normal notes; holds sustain long notes; a double press triggers a marked accent.",
            "Team notes require two or three named pods within the timing window.",
            "The hub scores Perfect, Good, Late, Early, or Miss using calibrated event timestamps.",
            "A clean streak adds a gentle LED trail; misses use a brief neutral pulse rather than punitive noise."
        ],
        "rules": [
            "Only press when your pod cues unless the track shows an all-player unison symbol.",
            "A player cannot gain extra points from repeated presses inside one note window.",
            "Each track contains solo notes, relay notes that move around the group, and one team finish.",
            "Team score combines timing accuracy and synchronization; one very early press can break a team note.",
            "Highest Groove after the chosen set wins. Cooperative mode wins by reaching the target accuracy."
        ],
        "strategy": "Watch the pre-cue without anticipating the actual beat. In team mode, count aloud during practice, then rely on the track. Accurate simple notes are worth more than frantic extra presses.",
        "variations": [
            "Silent rhythm: vibration and LED only.",
            "Conductor: one player designs a short live sequence and the group repeats it.",
            "Ghabga round: slower call-and-response patterns with room for conversation between tracks."
        ],
        "safety": "Keep volume comfortable and provide visual/vibration alternatives. Pods stay in hands or on the table; no striking furniture, bodies, or devices to create sound."
    },
    {
        "number": 7,
        "name": "Switchback",
        "tagline": "Route a live signal while the meaning of every press keeps changing.",
        "tags": ["LOGIC", "FAST DECISIONS", "SEATED"],
        "meta": [
            ("Players", "4-6"), ("Duration", "8-12 minutes"),
            ("Device position", "Handheld in a circle"), ("Activity", "Low"),
            ("Mode", "Competitive or cooperative"), ("Best space", "Table, majlis, or circle")
        ],
        "objective": "Keep the digital Signal moving to valid players while global Switchback rules reverse or modify short, long, and double presses.",
        "setup": [
            "Seat or stand players in a clear circle and assign each a stable position number and pod color.",
            "Teach the base controls: short press passes clockwise, long hold passes counter-clockwise, double press jumps one player.",
            "Choose Cooperative survival or Competitive three-life mode.",
            "Start at slow speed with one modifier; increase speed only after a clean practice round."
        ],
        "device": [
            "The active pod pulses white and vibrates once; that player has 2.5 seconds to act.",
            "The hub announces active modifiers on screen and with a unique sound/vibration pattern.",
            "Reverse swaps short and long directions. Mirror requires the next player to repeat the same gesture. Color Gate permits passing only to the displayed color group.",
            "Lockdown names a target; the group must route the Signal to that player within three moves.",
            "The hub rejects invalid destinations and records a Fault without moving the Signal."
        ],
        "rules": [
            "Only the active player may issue a route command.",
            "A valid command moves the Signal immediately and resets the decision timer.",
            "Every 20-35 seconds the hub adds, removes, or swaps one modifier; no more than two modifiers are active in Standard mode.",
            "A timeout, invalid gesture, or impossible destination produces one Fault.",
            "Competitive mode removes one life from the player who caused the Fault; at zero they remain in the circle as a pass-through position but cannot win.",
            "Cooperative mode wins by surviving eight minutes with fewer than ten group Faults."
        ],
        "strategy": "Say the current rule aloud when it changes, but do not shout conflicting instructions at the active player. Plan one move ahead during Lockdown and remember that the shortest route may become invalid under Color Gate.",
        "variations": [
            "Calm mode: four-second decisions and one modifier at a time.",
            "Memory mode: the modifier disappears from the screen after five seconds.",
            "Team lanes: odd and even positions score separately while sharing the same Signal."
        ],
        "safety": "Play seated or standing in place. The challenge is cognitive, so physical speed is unnecessary; do not throw, slap, or pass the hardware itself."
    },
    {
        "number": 8,
        "name": "Ghost Frequency",
        "tagline": "Run pair tests, compare evidence, and identify the hidden Drifter.",
        "tags": ["DEDUCTION", "EVIDENCE", "LOW MOVEMENT"],
        "meta": [
            ("Players", "4-6"), ("Duration", "12-15 minutes"),
            ("Device position", "Handheld or table"), ("Activity", "Low"),
            ("Mode", "Group vs hidden Drifter"), ("Best space", "Table or discussion circle")
        ],
        "objective": "Operators must identify the hidden Drifter by running limited pair tests. The Drifter wins by surviving three investigation rounds and may corrupt one result.",
        "setup": [
            "Assign every player a pod and ask players to look away while roles are delivered by vibration.",
            "One player receives the Drifter pattern; all others receive the Operator pattern.",
            "The hub secretly assigns each pod a frequency group while ensuring the evidence has one unique solution.",
            "Give the group six pair tests across three investigation rounds and the Drifter one Interference charge."
        ],
        "device": [
            "A player opens a pair test by pressing another player's pod once; both players then hold their own pods for two seconds.",
            "The hub returns SAME or DIFFERENT through green/amber LED patterns and logs the result on the public evidence board.",
            "The Drifter may double-press during one of their tests to spend Interference and invert that result; the public log marks that one result in the game may be corrupted.",
            "During voting, each pod cycles through candidate colors with short presses. A long hold commits the vote and turns the LED off to hide the choice.",
            "The hub reveals all votes together after every active player has committed."
        ],
        "rules": [
            "Round one permits two tests, round two permits two, and round three permits the final two.",
            "The same pair cannot be tested twice unless a special Recheck token is earned by a unanimous discussion vote.",
            "Players may share reasoning, bluff, or accuse, but may not reveal the private role vibration by replaying or imitating it physically.",
            "A majority accusation at the end of a round removes that player from testing but not from discussion.",
            "Operators win immediately if the Drifter is removed. The Drifter wins after the third vote or if only two active players remain.",
            "If the group finds a contradiction, they must decide which result was corrupted rather than assuming the system failed."
        ],
        "strategy": "Build a simple evidence graph: SAME links likely share a group; DIFFERENT links separate groups. The Drifter should corrupt a result that creates two plausible stories, not an obvious contradiction.",
        "variations": [
            "Open-data mode: no corruption charge; ideal for first-time players.",
            "Two Drifters: six experienced players, eight tests, and one shared Interference charge.",
            "Cooperative logic mode: no hidden role; identify which pod the hub intentionally configured differently."
        ],
        "safety": "Keep the discussion playful. Ban personal insults, real-world accusations, exclusion outside the game, and pressure to reveal private information. The hidden role is fictional and ends with the round."
    },
]


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def pil_font(size, bold=False):
    f = r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"
    return ImageFont.truetype(f, size)


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color is not None: run.font.color.rgb = rgb(color)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tcMar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}"); tcMar.append(node)
        node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")


def mark_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
    tblW.set(qn("w:w"), "9360"); tblW.set(qn("w:type"), "dxa")
    ind = tblPr.find(qn("w:tblInd"))
    if ind is None:
        ind = OxmlElement("w:tblInd"); tblPr.append(ind)
    ind.set(qn("w:w"), "120"); ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(round(width * 1440))); grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tcW = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            tcW.set(qn("w:w"), str(round(width * 1440))); tcW.set(qn("w:type"), "dxa")


def add_picture(run, path, alt, width=None, height=None):
    shape = run.add_picture(str(path), width=width, height=height)
    shape._inline.docPr.set("descr", alt)
    shape._inline.docPr.set("title", alt)
    return shape


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page "); set_run_font(run, size=8.5, color=GRAY)
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE"); paragraph._p.append(fld)


def restart_numbering(doc, items):
    numbering = doc.part.numbering_part.element
    style = doc.styles["List Number"]._element
    base_id = style.pPr.numPr.numId.val
    base = next(n for n in numbering.findall(qn("w:num")) if int(n.get(qn("w:numId"))) == base_id)
    abstract_id = base.find(qn("w:abstractNumId")).get(qn("w:val"))
    new_id = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(new_id))
    abstract = OxmlElement("w:abstractNumId"); abstract.set(qn("w:val"), abstract_id); num.append(abstract)
    override = OxmlElement("w:lvlOverride"); override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride"); start.set(qn("w:val"), "1"); override.append(start); num.append(override)
    numbering.append(num)
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)
        numPr = p._p.get_or_add_pPr().get_or_add_numPr()
        numPr.get_or_add_ilvl().val = 0; numPr.get_or_add_numId().val = new_id


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_label_para(doc, label, text):
    p = doc.add_paragraph()
    r = p.add_run(label + " "); r.bold = True; r.font.color.rgb = rgb(NAVY)
    p.add_run(text)
    return p


def add_meta_table(doc, meta):
    table = doc.add_table(rows=3, cols=4)
    table.style = "Table Grid"
    mark_header(table.rows[0])
    for i, (label, value) in enumerate(meta):
        row = i // 2; pair = i % 2
        c1 = table.cell(row, pair * 2); c2 = table.cell(row, pair * 2 + 1)
        c1.text = label; c2.text = value
        shade_cell(c1, TABLE_FILL)
        for run in c1.paragraphs[0].runs:
            set_run_font(run, size=9, bold=True, color=NAVY)
        for run in c2.paragraphs[0].runs:
            set_run_font(run, size=9, color=NAVY)
    set_table_geometry(table, [1.0, 2.25, 1.0, 2.25])
    return table


def draw_banner(game):
    path = ASSETS / f"game-{game['number']:02d}.png"
    img = Image.new("RGB", (1800, 520), "#F4F7FA")
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((40, 40, 1760, 480), 35, fill="white", outline="#D4DEE7", width=4)
    d.rounded_rectangle((70, 70, 390, 450), 30, fill="#102A43")
    d.text((120, 105), f"{game['number']:02d}", font=pil_font(125, True), fill="#7FDBD4")
    x0, y0 = 230, 330
    n = game["number"]
    if n == 1:
        d.polygon([(130,370),(165,300),(210,350),(255,295),(300,370)], fill="#E9A23B")
        d.rectangle((140,370,290,392), fill="#E9A23B")
    elif n == 2:
        d.line((120,345,305,345), fill="#E76F51", width=18)
        d.polygon([(305,345),(265,320),(265,370)], fill="#E76F51")
        d.ellipse((110,315,170,375), outline="#7FDBD4", width=10)
    elif n == 3:
        d.rounded_rectangle((125,320,205,405), 12, outline="#E9A23B", width=12)
        d.rounded_rectangle((225,320,305,405), 12, outline="#7FDBD4", width=12)
        d.arc((145,270,285,360), 180, 360, fill="white", width=12)
    elif n == 4:
        d.line((125,320,300,320), fill="#7FDBD4", width=14)
        d.polygon([(300,320),(270,300),(270,340)], fill="#7FDBD4")
        d.line((300,385,125,385), fill="#E76F51", width=14)
        d.polygon([(125,385),(155,365),(155,405)], fill="#E76F51")
    elif n == 5:
        for yy in (300,370):
            for xx in (140,220,300): d.ellipse((xx-18,yy-18,xx+18,yy+18), fill="#7FDBD4")
        for yy in (300,370): d.line((140,yy,300,yy), fill="#E9A23B", width=8)
        for xx in (140,220,300): d.line((xx,300,xx,370), fill="#E9A23B", width=8)
    elif n == 6:
        pts=[]
        for x in range(120,320,8): pts.append((x,345+int(45*math.sin((x-120)/18))))
        d.line(pts, fill="#E76F51", width=10)
    elif n == 7:
        d.arc((125,285,305,420), 205, 520, fill="#7FDBD4", width=14)
        d.polygon([(125,350),(155,325),(155,372)], fill="#7FDBD4")
        d.polygon([(305,350),(275,325),(275,372)], fill="#E9A23B")
    else:
        for r, color in ((85,"#7FDBD4"),(55,"#E9A23B"),(25,"#E76F51")):
            d.ellipse((215-r,350-r,215+r,350+r), outline=color, width=9)
    d.text((470, 105), game["name"], font=pil_font(62, True), fill="#102A43")
    d.text((470, 190), game["tagline"], font=pil_font(34), fill="#5C6770")
    x = 470
    for tag in game["tags"]:
        w = d.textlength(tag, font=pil_font(23, True)) + 46
        d.rounded_rectangle((x, 300, x+w, 360), 18, fill="#E4F4F4", outline="#0C7C86", width=2)
        d.text((x+23, 316), tag, font=pil_font(23, True), fill="#0C7C86")
        x += w + 20
    img.save(path, quality=95)
    return path


def draw_library_map():
    path = ASSETS / "library-map.png"
    img = Image.new("RGB", (1800, 1080), "white")
    d = ImageDraw.Draw(img)
    d.text((90, 70), "EIGHT NEW WAYS TO USE THE SAME SIX PODS", font=pil_font(50, True), fill="#102A43")
    d.text((90, 145), "No new sensors. Each game changes positioning, timing, feedback, and rules.", font=pil_font(29), fill="#5C6770")
    for i, game in enumerate(GAMES):
        col=i%2; row=i//2
        x=90+col*840; y=240+row*190
        accent=["#E9A23B","#E76F51","#0C7C86","#2E74B5"][row]
        d.rounded_rectangle((x,y,x+780,y+145),25,fill="#F4F7FA",outline=accent,width=4)
        d.ellipse((x+25,y+35,x+100,y+110),fill=accent)
        d.text((x+50,y+50),str(i+1),font=pil_font(30,True),fill="white",anchor="mm")
        d.text((x+125,y+28),game["name"],font=pil_font(30,True),fill="#102A43")
        d.text((x+125,y+75),game["tags"][0]+" / "+game["tags"][1],font=pil_font(23),fill="#5C6770")
    img.save(path, quality=95)
    return path


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = Inches(1); section.bottom_margin = Inches(1)
    section.left_margin = Inches(1); section.right_margin = Inches(1)
    section.header_distance = Inches(.492); section.footer_distance = Inches(.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(11); normal.font.color.rgb = rgb("253746")
    normal.paragraph_format.space_before = Pt(0); normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in [
        ("Heading 1",16,BLUE,18,10),("Heading 2",13,BLUE,14,7),("Heading 3",12,DARK_BLUE,10,5)
    ]:
        st=doc.styles[name]; st.font.name="Calibri"; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=rgb(color)
        st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True
    for name in ("List Bullet","List Number"):
        st=doc.styles[name]; st.font.name="Calibri"; st.font.size=Pt(11)
        st.paragraph_format.left_indent=Inches(.375); st.paragraph_format.first_line_indent=Inches(-.188)
        st.paragraph_format.space_after=Pt(4); st.paragraph_format.line_spacing=1.25
    return section


def add_header_footer(section):
    p = section.header.paragraphs[0]
    p.text = "PHYSICAL SOCIAL GAMES  |  EXPANSION PACK"
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in p.runs: set_run_font(run,size=8,bold=True,color=GRAY)
    add_page_field(section.footer.paragraphs[0])


def add_cover(doc):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(14)
    r=p.add_run("GAME LIBRARY  •  EXPANSION PACK 01"); set_run_font(r,size=10,bold=True,color=TEAL)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(7)
    r=p.add_run("Eight New Games"); set_run_font(r,size=30,bold=True,color=NAVY)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(6)
    r=p.add_run("Built for the Same Six Wearable Devices and Hub"); set_run_font(r,size=15,bold=True,color=DARK_BLUE)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(18)
    r=p.add_run("Complete objectives, setup, rules, device behavior, winning conditions, variations, and safety notes"); set_run_font(r,size=10.5,color=GRAY)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_picture(p.add_run(),KIT_IMAGE,"Concept image of six wearable game pods, one hub, and a carry case",width=Inches(6.45))
    c=doc.add_paragraph("One physical system. Eight additional styles of play."); c.alignment=WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(c.runs[0],size=9,italic=True,color=GRAY)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(10)
    r=p.add_run("Working game names - August 2026"); set_run_font(r,size=9,color=GRAY)
    doc.add_page_break()


def add_overview(doc, library_map):
    doc.add_heading("Expansion pack overview", level=1)
    doc.add_paragraph("These eight games are additional concepts. They do not repeat the first library's chase elimination, point theft, memory chain, guardian rescue, territory capture, pattern copy, background mission, or relay mechanics. Every concept uses the same six press-capable pods, RGB feedback, vibration, optional sound, and central hub described in the main business handbook.")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_picture(p.add_run(),library_map,"Map of eight new games and their primary mechanics",width=Inches(6.4))
    c=doc.add_paragraph("The hardware stays fixed; device placement and software rules create the variety."); c.alignment=WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(c.runs[0],size=9,italic=True,color=GRAY)
    doc.add_heading("Hardware modes used",level=2)
    modes=[
        ("Wearable/front clip","Players press another person's pod or their own safely."),
        ("Handheld","Private haptics, rhythm, voting, trading, and logic inputs."),
        ("Table/station","Pods become generators, locks, terminals, or fixed controls."),
        ("Hub screen","Public timer, state, clues, evidence, market, scoring, and winner."),
    ]
    table=doc.add_table(rows=1,cols=2); table.style="Table Grid"; table.cell(0,0).text="Mode"; table.cell(0,1).text="Use"
    shade_cell(table.cell(0,0),TABLE_FILL); shade_cell(table.cell(0,1),TABLE_FILL); mark_header(table.rows[0])
    for mode,use in modes:
        row=table.add_row().cells; row[0].text=mode; row[1].text=use
    set_table_geometry(table,[1.7,4.8])
    for i,row in enumerate(table.rows):
        for cell in row.cells:
            for run in cell.paragraphs[0].runs: set_run_font(run,size=9.5,bold=(i==0),color=NAVY)
    doc.add_heading("Library at a glance",level=2)
    table=doc.add_table(rows=1,cols=4); table.style="Table Grid"
    for i,t in enumerate(("Game","Primary mechanic","Players","Activity")):
        table.cell(0,i).text=t; shade_cell(table.cell(0,i),TABLE_FILL)
    mark_header(table.rows[0])
    for game in GAMES:
        row=table.add_row().cells
        row[0].text=game["name"]; row[1].text=game["tags"][0].title(); row[2].text=game["meta"][0][1]; row[3].text=game["meta"][3][1]
    set_table_geometry(table,[1.55,2.45,.8,1.7])
    for i,row in enumerate(table.rows):
        for cell in row.cells:
            for run in cell.paragraphs[0].runs: set_run_font(run,size=8.8,bold=(i==0),color=NAVY)
    doc.add_heading("Universal safety rule",level=2)
    add_label_para(doc,"Applies to every game:","touch only the device; never push, grab, tackle, restrain, block doors, or play near roads, stairs, pools, kitchens, vehicles, breakables, or other hazards. Stop for damaged hardware, heat, pain, unsafe behavior, or a loose attachment. The prototype remains supervised and 14+ until product safety and age suitability are formally validated.")
    doc.add_page_break()


def add_game(doc, game, banner):
    p=doc.add_paragraph(); p.paragraph_format.page_break_before=True; p.paragraph_format.space_after=Pt(8)
    add_picture(p.add_run(),banner,f"Illustrated title banner for game {game['number']}: {game['name']}",width=Inches(6.5))
    add_meta_table(doc,game["meta"])
    doc.add_heading("Objective",level=2)
    doc.add_paragraph(game["objective"])
    doc.add_heading("Setup",level=2)
    restart_numbering(doc,game["setup"])
    doc.add_heading("How the pods and hub behave",level=2)
    add_bullets(doc,game["device"])
    doc.add_heading("Rules and round flow",level=2)
    restart_numbering(doc,game["rules"])
    doc.add_heading("Strategy and host guidance",level=2)
    doc.add_paragraph(game["strategy"])
    doc.add_heading("Variations",level=2)
    add_bullets(doc,game["variations"])
    doc.add_heading("Safety note",level=2)
    add_label_para(doc,"Host rule:",game["safety"])


def add_closing(doc):
    p=doc.add_paragraph(); p.paragraph_format.page_break_before=True
    doc.add_heading("Playtest and release checklist",level=1)
    doc.add_paragraph("A concept enters the sellable library only after its rules, software state, physical behavior, and recovery cases are tested. Use the same evidence standard for every game.")
    items=[
        "Run one paper or facilitator-led version before implementing software.",
        "Test every short, long, double, simultaneous, early, late, duplicate, and disconnected input that the game uses.",
        "Define what happens when a player or pod leaves during the round.",
        "Measure explanation time, setup time, round length, downtime, false inputs, and requests for another round.",
        "Record safety interventions and remove any rule that rewards unsafe contact or movement.",
        "Confirm the game feels meaningfully different from the existing library.",
        "Lock the rule version, game-engine schema, feedback patterns, and host instructions together.",
        "Ship only after unfamiliar hosts can run it from the guide without developer help."
    ]
    restart_numbering(doc,items)
    doc.add_heading("Recommended implementation order",level=2)
    doc.add_paragraph("Start with the least firmware risk and increase complexity gradually:")
    restart_numbering(doc,[
        "Double Lock - validates simultaneous holds and station mode.",
        "Switchback - validates short, long, and double gesture routing.",
        "Power Grid - validates continuous state and timed events.",
        "Rhythm Rivals - validates calibrated timing and group synchronization.",
        "Crown Circuit - validates duel locking and high-tempo acknowledgements.",
        "Market Rush - validates transactional UI and two-party confirmation.",
        "Courier Code - validates hidden roles and private haptics.",
        "Ghost Frequency - validates evidence generation, secret voting, and rule-consistent deception."
    ])
    doc.add_page_break()
    kicker=doc.add_paragraph(); kicker.alignment=WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_before=Pt(26); kicker.paragraph_format.space_after=Pt(4)
    kr=kicker.add_run("EXPANSION PACK 01")
    set_run_font(kr,size=9,bold=True,color=BLUE)
    title=doc.add_paragraph(); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after=Pt(14)
    tr=title.add_run("Eight more reasons to replay the same box.")
    set_run_font(tr,size=20,bold=True,color=NAVY)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_picture(p.add_run(),PLAY_IMAGE,"Six adults playing safely with wearable game pods in a cleared villa garden",width=Inches(6.25))
    c=doc.add_paragraph("One six-pod kit. Eight distinct new social experiences. No extra hardware."); c.alignment=WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(c.runs[0],size=10,bold=True,color=NAVY)
    n=doc.add_paragraph("Build first: Double Lock → Switchback → Power Grid")
    n.alignment=WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(n.runs[0],size=9,italic=True,color=GRAY)


def build():
    banners=[draw_banner(g) for g in GAMES]
    library_map=draw_library_map()
    doc=Document(); section=configure_styles(doc); add_header_footer(section)
    add_cover(doc); add_overview(doc,library_map)
    for game,banner in zip(GAMES,banners): add_game(doc,game,banner)
    add_closing(doc)
    doc.core_properties.title="Eight New Games Expansion Pack"
    doc.core_properties.subject="Eight original games using the same six wearable electronic pods and hub"
    doc.core_properties.author="Physical Social Games Project"
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    build()

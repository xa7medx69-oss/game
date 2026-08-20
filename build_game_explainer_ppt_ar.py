from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from build_game_explainer_ppt import (
    AMBER,
    ASSET_DIR,
    BLUE,
    CORAL,
    CREAM,
    GRAY,
    HANDBOOK_ASSETS,
    INK,
    MINT,
    NAVY,
    PALE,
    PEACH,
    SKY,
    SLIDE_H,
    SLIDE_W,
    TEAL,
    WHITE,
    add_picture_cover,
    add_round_rect,
    rgb,
)


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Physical-Social-Games-Kids-Explainer-Arabic.pptx"
ARABIC_FONT = "Dubai"


GAMES = [
    {
        "name": "صيد الإشارة",
        "tagline": "اهرب من الصياد حتى ينتهي الوقت",
        "players": "٤–٦",
        "movement": "حركة كثيرة",
        "image": "01-pulse-hunt.png",
        "goal": "لاعب واحد هو الصياد، والباقون يهربون منه حتى ينتهي الوقت.",
        "action": "يلمس الصياد الزر المثبّت على ظهر اللاعب. إذا أضاء الزر، يخرج اللاعب ويتجه إلى المكان الآمن.",
        "win": "يفوز الهاربون إذا بقي واحد منهم في اللعب. ويفوز الصياد إذا أخرج الجميع.",
        "safety": "نلمس الزر فقط، ولا ندفع أحدًا ولا نمسكه أو نغلق الطريق أمامه.",
    },
    {
        "name": "سرقة الجواهر",
        "tagline": "اسرق من الفريق الآخر واحمِ جواهرك",
        "players": "٤–٦",
        "movement": "حركة متوسطة",
        "image": "02-vault-heist.png",
        "goal": "لكل فريق خزنة مليئة بجواهر رقمية. حاول جمع جواهر أكثر من الفريق الآخر.",
        "action": "لسرقة جوهرة، اضغط مطولًا على زر لاعب من الفريق الآخر. بعدها يحصل على درع مؤقت.",
        "win": "عندما تنتهي الجولة، يفوز الفريق الذي بقيت لديه جواهر أكثر.",
        "safety": "اضغط الزر بلطف. لا تمسك اللاعب أو ملابسه، ولا تمنعه من الحركة.",
    },
    {
        "name": "سلسلة الأضواء",
        "tagline": "احفظ ترتيب الأضواء ولا تقطع السلسلة",
        "players": "٣–٦",
        "movement": "حركة خفيفة",
        "image": "03-signal-chain.png",
        "goal": "تضيء أزرار اللاعبين واحدًا بعد الآخر. مهمتكم أن تتذكروا الترتيب الصحيح.",
        "action": "بعد أن تشاهدوا الترتيب، يضغط كل لاعب زره عندما يأتي دوره.",
        "win": "إذا كررتم السلسلة كاملة قبل انتهاء الوقت، يفوز الفريق كله.",
        "safety": "ابقَ في مكانك وانتظر دورك؛ هذه لعبة ذاكرة وليست لعبة جري.",
    },
    {
        "name": "الشريك الحارس",
        "tagline": "اعرف شريكك سرًا وأنقذه في الوقت المناسب",
        "players": "٤–٦",
        "movement": "حركة متوسطة",
        "image": "04-guardian-link.png",
        "goal": "لكل لاعب شريك سري. أنجزوا المهام وحاولوا حماية بعضكم من الخروج.",
        "action": "إذا حاول أحد إخراج شريكك، اضغط زر شريكك خلال ثلاث ثوانٍ لتمنحه فرصة نجاة.",
        "win": "يفوز الثنائي الذي يحمي بعضه وينجز أكبر عدد من المهام.",
        "safety": "الإنقاذ يكون بلمس الزر فقط. لا تلمس جسم اللاعب ولا تقف في طريقه.",
    },
    {
        "name": "السيطرة على المناطق",
        "tagline": "اجعل أكبر عدد من المناطق بلون فريقك",
        "players": "٤–٦",
        "movement": "حركة متوسطة",
        "image": "05-territory-pulse.png",
        "goal": "توجد عدة مناطق في المكان. كل فريق يحاول تحويلها إلى لون فريقه.",
        "action": "اضغط مطولًا على زر المنطقة حتى يتغير لونه. تبدأ المنطقة بعدها بجمع النقاط لفريقك.",
        "win": "يفوز الفريق الذي يحافظ على مناطق أكثر لمدة أطول.",
        "safety": "امشِ بين المناطق، واترك الأجهزة في مكان واضح حتى لا يدوس عليها أحد.",
    },
    {
        "name": "قلّد النمط",
        "tagline": "شاهد الألوان ثم كررها بالترتيب",
        "players": "٢–٦",
        "movement": "حركة خفيفة",
        "image": "06-echo-match.png",
        "goal": "تعرض الشاشة نمطًا قصيرًا من الألوان أو الأصوات، وعليك أن تتذكره.",
        "action": "عندما ينتهي العرض، اضغط زرّك لتكرر النمط بالترتيب نفسه.",
        "win": "كل إجابة صحيحة تكسبك نقاطًا. يفوز اللاعب الأسرع والأدق.",
        "safety": "ضع الجهاز بشكل ثابت، وخفّض الصوت إذا كان مرتفعًا أو مزعجًا.",
    },
    {
        "name": "المهمة السرية",
        "tagline": "أنجز مهمة بسيطة من دون أن يعرف الآخرون",
        "players": "٣–٦",
        "movement": "حركة خفيفة",
        "image": "07-silent-mission.png",
        "goal": "يحصل كل لاعب على مهمة خفية، مثل جعل شخص يضحك أو تحية أحد اللاعبين.",
        "action": "عندما تنجز جزءًا من مهمتك، سجّله بضغطة قصيرة أو طويلة على زرّك.",
        "win": "في نهاية الجولة، تكشف الشاشة من نجح في إكمال مهمته.",
        "safety": "يجب أن تكون كل مهمة لطيفة وآمنة، وألا تسبب الإحراج أو الإزعاج لأي شخص.",
    },
    {
        "name": "سباق الإشارة",
        "tagline": "مرّر الضوء من لاعب إلى آخر بأسرع وقت",
        "players": "٣–٦",
        "movement": "حركة خفيفة",
        "image": "08-relay-rush.png",
        "goal": "تبدأ إشارة مضيئة عند لاعب، ويجب أن تمر عبر الفريق حتى تصل إلى النهاية.",
        "action": "عندما يضيء زرّك، اضغطه فورًا لترسل الإشارة إلى اللاعب التالي.",
        "win": "يفوز الفريق إذا أكمل المسار بسرعة وبأقل عدد من الأخطاء.",
        "safety": "مرّر الإشارة فقط؛ لا ترمِ الجهاز ولا تمرره من يد إلى يد.",
    },
    {
        "name": "ملك التاج",
        "tagline": "احتفظ بالتاج واربح تحديات السرعة",
        "players": "٤–٦",
        "movement": "حركة متوسطة",
        "image": "09-crown-circuit.png",
        "goal": "أحد اللاعبين يحمل تاجًا افتراضيًا ويجمع نقاطًا ما دام التاج معه.",
        "action": "لتتحداه، اضغط زر التاج ثم ابتعدا. بعد ظهور الوميض، يتسابق كل منكما للضغط على زرّه.",
        "win": "من يحتفظ بالتاج ويفوز بالتحديات يصل إلى عدد النقاط المطلوب أولًا.",
        "safety": "بعد بدء التحدي ابتعد خطوة بهدوء. لا تدفع حامل التاج ولا تلاحقه بسرعة.",
    },
    {
        "name": "الرسالة السرية",
        "tagline": "مرّر الرسالة من دون أن يكتشفها المراقب",
        "players": "٥–٦",
        "movement": "لعبة جلوس",
        "image": "10-courier-code.png",
        "goal": "الفريق يحاول تمرير رسالة سرية، بينما يحاول لاعب خفي اكتشاف من يحملها.",
        "action": "يمسك حامل الرسالة زر زميله، ثم يضغط الزميل زرّه ليستلم الرسالة. يمكن للآخرين التظاهر بالتسليم.",
        "win": "يفوز الفريق إذا مرّر الرسالة ثلاث مرات وأنهى المهمة. ويفوز المراقب إذا اكتشف حاملها.",
        "safety": "لا تفتش ملابس أحد، ولا تمسك يده، ولا تجبره على كشف الإشارة التي شعر بها.",
    },
    {
        "name": "لغز القفل",
        "tagline": "حلوا الدليل واضغطوا الأزرار معًا",
        "players": "٣–٦",
        "movement": "لعبة طاولة",
        "image": "11-double-lock.png",
        "goal": "تعرض الشاشة لغزًا يقودكم إلى لونين أو ثلاثة ألوان صحيحة.",
        "action": "ناقشوا الحل، ثم اضغطوا الأزرار الصحيحة في اللحظة نفسها لمدة ثانيتين.",
        "win": "افتحوا الأقفال الستة قبل انتهاء الوقت لتفوزوا جميعًا.",
        "safety": "ضعوا الأزرار على طاولة ثابتة، واتركوا لكل لاعب مساحة مريحة للوصول.",
    },
    {
        "name": "سوق المقايضة",
        "tagline": "تفاوض مع الآخرين واجمع ما تحتاجه",
        "players": "٤–٦",
        "movement": "لعبة جلوس",
        "image": "12-market-rush.png",
        "goal": "لديك موارد خيالية: شرارة وقماش ووقت. تحتاجها لإكمال بطاقات الطلب.",
        "action": "اتفق مع لاعب على صفقة، وحددا ما سيعطيه كل منكما. تتم الصفقة فقط عندما توافقان معًا.",
        "win": "كل بطاقة طلب مكتملة تمنحك نجوم سمعة. يفوز صاحب أكبر عدد من النجوم.",
        "safety": "هذه موارد خيالية للعبة فقط؛ لا نستخدم مالًا حقيقيًا ولا مراهنات أو وعودًا.",
    },
    {
        "name": "أنقذ المدينة",
        "tagline": "وازن الطاقة ولا تدع المولدات تسخن",
        "players": "٣–٦",
        "movement": "لعبة طاولة",
        "image": "13-power-grid.png",
        "goal": "ستة مولدات تشغّل مدينة صغيرة. المطلوب منكم إنتاج كمية الطاقة التي تحتاجها المدينة.",
        "action": "ضغطة قصيرة تزيد الطاقة، وضغطة طويلة تبرد المولد. أحيانًا يجب إصلاح مولدين معًا.",
        "win": "إذا بقيت المدينة مستقرة ومضيئة حتى نهاية الوقت، يفوز الفريق كله.",
        "safety": "اترك الأجهزة مسطحة وجافة على الطاولة، ولا تضعها قرب المشروبات.",
    },
    {
        "name": "تحدّي الإيقاع",
        "tagline": "راقب الضوء واضغط مع النغمة",
        "players": "٢–٦",
        "movement": "حركة خفيفة",
        "image": "14-rhythm-rivals.png",
        "goal": "لكل لاعب نغمات خاصة به، والمطلوب الضغط في اللحظة الصحيحة.",
        "action": "يضيء زرّك قبل النغمة بقليل. انتظر الإشارة القوية ثم اضغط. بعض النغمات تحتاج الفريق معًا.",
        "win": "الضغط الدقيق يبني سلسلة نقاط. يفوز أعلى لاعب، أو يتعاون الجميع للوصول إلى هدف واحد.",
        "safety": "اجعل الصوت مريحًا، ولا تضرب الجهاز بالطاولة أو بأي شخص.",
    },
    {
        "name": "غيّر الاتجاه",
        "tagline": "مرّر الإشارة وانتبه لأن القاعدة قد تتغير",
        "players": "٤–٦",
        "movement": "لعبة جلوس",
        "image": "15-switchback.png",
        "goal": "إشارة مضيئة تنتقل حول الدائرة، ويجب أن تصل في كل مرة إلى لاعب صحيح.",
        "action": "الضغطة القصيرة ترسلها في اتجاه، والطويلة تعكس الاتجاه، والمزدوجة تقفز فوق لاعب. ثم تتغير القاعدة فجأة.",
        "win": "أكملوا الجولة بأقل عدد من الأخطاء، أو احتفظ بفرص أكثر من بقية اللاعبين.",
        "safety": "العبوا وأنتم جالسون. اضغطوا الأزرار فقط ولا ترموا الأجهزة أو تتبادلوها.",
    },
    {
        "name": "اكتشف المتخفي",
        "tagline": "اختبر اللاعبين واجمع الأدلة ثم صوّت",
        "players": "٤–٦",
        "movement": "لعبة جلوس",
        "image": "16-ghost-frequency.png",
        "goal": "أحد اللاعبين هو المتخفي، والباقون يحاولون اكتشافه باستخدام عدد قليل من الاختبارات.",
        "action": "اختبروا زرين معًا لتعرفوا هل هما متشابهان أم مختلفان. المتخفي يستطيع تغيير نتيجة واحدة لخداعكم.",
        "win": "يفوز الفريق إذا صوّت للشخص الصحيح، ويفوز المتخفي إذا بقي مجهولًا حتى النهاية.",
        "safety": "الشك والاتهام جزء من اللعبة فقط. تكلموا بلطف ولا تستبعدوا أحدًا بعد انتهاء الجولة.",
    },
]


def arabic_number(number):
    return str(number).translate(str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩"))


def set_arabic_run(run, font, size, color, bold):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "ar-AE")
    rpr.set("altLang", "en-US")
    for tag in ("a:cs", "a:ea"):
        node = OxmlElement(tag)
        node.set("typeface", font)
        rpr.append(node)


def add_ar_text(slide, text, x, y, w, h, size=20, color=INK, bold=False,
                align=PP_ALIGN.RIGHT, valign=MSO_ANCHOR.TOP, margin=0.06,
                line_spacing=1.0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    frame._txBody.bodyPr.set("rtlCol", "1")
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.line_spacing = line_spacing
    paragraph._p.get_or_add_pPr().set("rtl", "1")
    run = paragraph.add_run()
    run.text = text
    set_arabic_run(run, ARABIC_FONT, size, color, bold)
    return box


def add_top_bar(slide, section, number=None):
    add_round_rect(slide, 0, 0, SLIDE_W, 0.62, NAVY, MSO_SHAPE.RECTANGLE)
    add_ar_text(slide, "ألعاب جماعية تفاعلية", 8.35, 0.14, 4.55, 0.28,
                11, WHITE, True, margin=0)
    label = section if number is None else f"{section}  /  {arabic_number(f'{number:02d}')}"
    add_ar_text(slide, label, 0.42, 0.14, 4.7, 0.28, 10.5, "7FDBD4", True,
                align=PP_ALIGN.LEFT, margin=0)


def add_pill(slide, text, x, y, w, fill=MINT, color=TEAL):
    add_round_rect(slide, x, y, w, 0.42, fill, line=color, line_width=0.8)
    add_ar_text(slide, text, x + 0.06, y + 0.075, w - 0.12, 0.24, 9.2, color, True,
                align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)


def add_info_card(slide, label, body, x, y, w, fill, accent):
    add_round_rect(slide, x, y, w, 1.0, fill)
    add_round_rect(slide, x + w - 0.10, y, 0.10, 1.0, accent, MSO_SHAPE.RECTANGLE)
    add_ar_text(slide, label, x + 0.18, y + 0.10, w - 0.39, 0.21, 9.4, accent, True, margin=0)
    add_ar_text(slide, body, x + 0.18, y + 0.33, w - 0.39, 0.59, 12.7, INK,
                margin=0, line_spacing=0.94)


def add_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(NAVY)
    add_round_rect(slide, 7.43, 0.55, 5.35, 6.4, "163954")
    add_ar_text(slide, "١٦ لعبة،\nومجموعة أجهزة واحدة!", 7.88, 1.0, 4.45, 1.45, 29, WHITE, True,
                line_spacing=0.9)
    add_ar_text(slide, "شرح بسيط بالصور: ماذا نفعل؟ وكيف نفوز؟", 7.88, 2.72, 4.45, 0.92,
                19, "CFE8E8")
    add_round_rect(slide, 8.63, 4.05, 3.7, 0.58, AMBER)
    add_ar_text(slide, "شاهد  •  اضغط  •  العب", 8.76, 4.17, 3.45, 0.28, 14,
                NAVY, True, align=PP_ALIGN.CENTER, margin=0)
    add_ar_text(slide, "شرح مبسّط يناسب الصغار\nملاحظة: النموذج الحالي للاستخدام تحت إشراف الكبار ولعمر ١٤ سنة فأكثر.",
                7.88, 5.25, 4.45, 0.82, 12, "AFC3D4")
    add_round_rect(slide, 0.55, 0.55, 6.58, 6.4, WHITE)
    add_picture_cover(slide, HANDBOOK_ASSETS / "kit-concept.png", 0.70, 0.72, 6.28, 5.45,
                      "صورة تصورية لستة أجهزة لعب ووحدة مركزية وحقيبة حمل")
    add_ar_text(slide, "ستة أزرار ذكية ووحدة تحكم… وكل لعبة لها فكرة مختلفة", 0.96, 6.26, 5.7, 0.38,
                16, NAVY, True, align=PP_ALIGN.CENTER, margin=0)


def add_how_it_works_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(PALE)
    add_top_bar(slide, "كيف تعمل؟")
    add_ar_text(slide, "كيف تصبح الأزرار نفسها ١٦ لعبة مختلفة؟", 4.0, 0.88, 8.75, 0.55, 27, NAVY, True)
    add_ar_text(slide, "أنت تضغط الزر، ووحدة التحكم تفهم الضغطة وتطبّق قواعد اللعبة.",
                0.72, 1.44, 12.0, 0.45, 16, GRAY)
    cards = [
        ("١", "جهّز الزر", "ارتدِه بطريقة آمنة، أمسكه بيدك، أو ضعه ثابتًا على الطاولة.", TEAL, MINT),
        ("٢", "انتبه للضوء والاهتزاز", "فالزر يخبرك متى يبدأ دورك وماذا يحدث أثناء اللعب.", CORAL, PEACH),
        ("٣", "اضغط بالطريقة المطلوبة", "قد تكون ضغطة سريعة أو طويلة أو مزدوجة، وقد تضغطون معًا.", AMBER, CREAM),
    ]
    for i, (num, title, body, accent, fill) in enumerate(cards):
        x = 8.96 - i * 4.18
        add_round_rect(slide, x, 2.18, 3.75, 3.75, WHITE, line="D4DEE7", line_width=1)
        add_round_rect(slide, x + 1.36, 2.48, 1.02, 1.02, accent, MSO_SHAPE.OVAL)
        add_ar_text(slide, num, x + 1.36, 2.66, 1.02, 0.45, 24, WHITE, True,
                    align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        add_ar_text(slide, title, x + 0.34, 3.82, 3.07, 0.42, 15, accent, True,
                    align=PP_ALIGN.CENTER, margin=0)
        add_ar_text(slide, body, x + 0.39, 4.40, 2.97, 0.90, 15.5, INK,
                    align=PP_ALIGN.CENTER, margin=0, line_spacing=1.02)
    add_round_rect(slide, 0.62, 6.36, 12.05, 0.6, NAVY)
    add_ar_text(slide, "وحدة التحكم تعرف من ضغط، وتحسب النقاط، ثم تشغّل الأضواء والاهتزاز.",
                0.95, 6.49, 11.4, 0.30, 14, WHITE, True,
                align=PP_ALIGN.CENTER, margin=0)


def add_safety_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(PALE)
    add_top_bar(slide, "السلامة أولًا")
    add_ar_text(slide, "يمكن لطفل عمره ١٠ سنوات فهم الفكرة، لكن اللعب له شروط سلامة",
                2.5, 0.86, 10.25, 0.58, 27, NAVY, True)
    add_picture_cover(slide, HANDBOOK_ASSETS / "gameplay-villa.png", 5.55, 1.62, 7.2, 4.63,
                      "بالغون يلعبون بأمان بالأجهزة في حديقة خالية من العوائق")
    rules = [
        ("العمر والإشراف", "هذا نموذج أولي لعمر ١٤ سنة فأكثر، ويجب أن يكون اللعب تحت إشراف شخص بالغ.", CORAL, PEACH),
        ("المس الزر فقط", "لا تدفع أحدًا، ولا تمسكه أو تعرقله، ولا تلمس وجهه أو رقبته.", TEAL, MINT),
        ("اختر مكانًا آمنًا", "العب في مساحة خالية، بعيدًا عن الطرق والسلالم والمسابح والسيارات والأشياء القابلة للكسر.", AMBER, CREAM),
        ("توقّف فورًا عند الخطر", "أوقف اللعب إذا شعرت بألم أو حرارة، أو تلف الجهاز، أو ارتخى الحزام، أو أصبح اللعب خشنًا.", BLUE, SKY),
    ]
    for i, (title, body, accent, fill) in enumerate(rules):
        y = 1.62 + i * 1.18
        add_round_rect(slide, 0.63, y, 4.65, 0.98, fill)
        add_ar_text(slide, title, 0.84, y + 0.10, 4.13, 0.23, 10, accent, True, margin=0)
        add_ar_text(slide, body, 0.84, y + 0.36, 4.13, 0.49, 12.7, INK,
                    margin=0, line_spacing=0.94)
    add_ar_text(slide, "قبل كل جولة: يفحص المشرف المكان، ويثبّت الأجهزة، ثم يشرح القواعد للجميع.", 0.75, 6.48, 11.9, 0.34,
                15, NAVY, True, align=PP_ALIGN.CENTER, margin=0)


def add_game_slide(prs, game, index):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(PALE)
    add_top_bar(slide, "مكتبة الألعاب", index)
    add_ar_text(slide, game["name"], 6.68, 0.82, 6.10, 0.55, 28, NAVY, True)
    add_ar_text(slide, game["tagline"], 6.38, 1.34, 6.40, 0.34, 14, GRAY)
    add_pill(slide, f"لاعبون: {game['players']}", 0.55, 0.88, 1.55)
    add_pill(slide, game["movement"], 2.24, 0.88, 1.42, CREAM, AMBER)
    add_pill(slide, "وحدة تحكم", 3.80, 0.88, 1.35, SKY, BLUE)

    add_round_rect(slide, 5.28, 1.82, 7.55, 4.78, WHITE, line="D4DEE7", line_width=1)
    add_picture_cover(slide, ASSET_DIR / game["image"], 5.39, 1.93, 7.33, 4.56,
                      f"شرح بصري للعبة {game['name']}")

    add_info_card(slide, "فكرة اللعبة", game["goal"], 0.55, 1.82, 4.48, MINT, TEAL)
    add_info_card(slide, "كيف نلعب؟", game["action"], 0.55, 3.02, 4.48, PEACH, CORAL)
    add_info_card(slide, "كيف نفوز؟", game["win"], 0.55, 4.22, 4.48, CREAM, AMBER)

    add_round_rect(slide, 0.55, 5.46, 4.48, 1.14, NAVY)
    add_ar_text(slide, "انتبه للسلامة", 3.42, 5.59, 1.32, 0.21, 9.5, "7FDBD4", True, margin=0)
    add_ar_text(slide, game["safety"], 0.80, 5.87, 3.93, 0.55, 12.5, WHITE,
                margin=0, line_spacing=0.93)
    add_ar_text(slide, arabic_number(f"{index:02d}"), 12.20, 6.82, 0.48, 0.23,
                9, GRAY, True, margin=0)
    add_ar_text(slide, "شرح مبسّط بالصور • التفاصيل الكاملة في دليل اللعبة", 0.62, 6.82, 4.55, 0.23,
                9, GRAY, align=PP_ALIGN.LEFT, margin=0)


def add_chooser_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(PALE)
    add_top_bar(slide, "اختر لعبتك")
    add_ar_text(slide, "ما نوع اللعبة التي تريدها اليوم؟", 3.0, 0.88, 9.75, 0.56, 27, NAVY, True)
    groups = [
        ("تحب الحركة؟", "صيد الإشارة\nسرقة الجواهر\nالشريك الحارس\nالسيطرة على المناطق\nملك التاج", CORAL, PEACH),
        ("تحب حل الألغاز؟", "سلسلة الأضواء\nسباق الإشارة\nلغز القفل\nأنقذ المدينة", TEAL, MINT),
        ("تحب الكلام والتمويه؟", "المهمة السرية\nالرسالة السرية\nسوق المقايضة\nاكتشف المتخفي", AMBER, CREAM),
        ("تحب السرعة والتركيز؟", "قلّد النمط\nتحدّي الإيقاع\nغيّر الاتجاه", BLUE, SKY),
    ]
    for i, (title, games, accent, fill) in enumerate(groups):
        x = 9.90 - i * 3.15
        add_round_rect(slide, x, 1.83, 2.83, 4.6, WHITE, line="D4DEE7", line_width=1)
        add_round_rect(slide, x, 1.83, 2.83, 0.72, accent, MSO_SHAPE.RECTANGLE)
        add_ar_text(slide, title, x + 0.15, 2.04, 2.53, 0.29, 13, WHITE, True,
                    align=PP_ALIGN.CENTER, margin=0)
        add_ar_text(slide, games, x + 0.30, 2.87, 2.23, 2.85, 17, NAVY, True,
                    align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0,
                    line_spacing=1.18)
        add_round_rect(slide, x + 0.96, 5.83, 0.90, 0.22, fill)
    add_round_rect(slide, 1.47, 6.70, 10.38, 0.45, NAVY)
    add_ar_text(slide, "اختر اللعبة التي تناسب المكان والوقت وطاقة المجموعة.",
                1.70, 6.80, 9.9, 0.24, 13, WHITE, True,
                align=PP_ALIGN.CENTER, margin=0)


def build():
    for game in GAMES:
        path = ASSET_DIR / game["image"]
        if not path.exists():
            raise FileNotFoundError(path)

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    props = prs.core_properties
    props.title = "ألعاب جماعية تفاعلية - شرح مبسّط بالصور"
    props.subject = "شرح عربي واضح وطبيعي لأفكار الألعاب الست عشرة"
    props.author = "مشروع الألعاب الجماعية التفاعلية"
    props.keywords = "ألعاب جماعية، ألعاب حركية، شرح بالصور، أزرار ذكية"
    props.comments = "شرح مبسّط للفهم. النموذج الأولي الحالي للاستخدام تحت الإشراف ولعمر ١٤ سنة فأكثر."

    add_title_slide(prs)
    add_how_it_works_slide(prs)
    add_safety_slide(prs)
    for index, game in enumerate(GAMES, 1):
        add_game_slide(prs, game, index)
    add_chooser_slide(prs)

    prs.save(OUTPUT)
    print(f"Saved {OUTPUT} with {len(prs.slides)} slides")


if __name__ == "__main__":
    build()

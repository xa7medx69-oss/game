from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
ASSET_DIR = ROOT / "presentation-assets" / "games"
HANDBOOK_ASSETS = ROOT / "physical-social-games-handbook" / "assets"
OUTPUT = ROOT / "Physical-Social-Games-Kids-Explainer.pptx"

SLIDE_W = 13.333
SLIDE_H = 7.5

NAVY = "102A43"
TEAL = "0C7C86"
CORAL = "E76F51"
AMBER = "E9A23B"
BLUE = "2E74B5"
INK = "253746"
GRAY = "5C6770"
PALE = "F4F7FA"
WHITE = "FFFFFF"
MINT = "E4F4F4"
PEACH = "FCEAE5"
CREAM = "FFF4DF"
SKY = "E8F1FA"


GAMES = [
    {
        "name": "Pulse Hunt",
        "tagline": "A safe chase-and-tag game",
        "players": "4-6",
        "movement": "Active",
        "image": "01-pulse-hunt.png",
        "goal": "Runners stay free until the timer ends. The Hunter tries to tag every Runner.",
        "action": "The Hunter presses only a Runner's back pod. A confirmed light means the tag counted.",
        "win": "Runners win if anyone is still free. The Hunter wins by tagging everyone.",
        "safety": "Touch only the pod - never push, grab, tackle, or block anyone.",
    },
    {
        "name": "Vault Heist",
        "tagline": "Steal gems, then protect your vault",
        "players": "4-6",
        "movement": "Medium",
        "image": "02-vault-heist.png",
        "goal": "Collect digital gems while stopping the other team from emptying your vault.",
        "action": "Hold an opponent's pod to steal one gem. A shield protects them for a short time.",
        "win": "Protect your vault and finish with more gems than the other team.",
        "safety": "Press the device gently. Do not hold, pull, or trap another player.",
    },
    {
        "name": "Signal Chain",
        "tagline": "Remember the order and keep the chain alive",
        "players": "3-6",
        "movement": "Low",
        "image": "03-signal-chain.png",
        "goal": "Work together to remember a growing order of players and colors.",
        "action": "Watch which pods light up. Then press your own pod when your turn arrives.",
        "win": "Complete the whole chain before the timer runs out.",
        "safety": "Stay in your place and wait for your turn - speed comes from thinking.",
    },
    {
        "name": "Guardian Link",
        "tagline": "Secret partners protect each other",
        "players": "4-6",
        "movement": "Medium",
        "image": "04-guardian-link.png",
        "goal": "Complete public tasks while secretly keeping your partner safe.",
        "action": "If your partner is hit, press their pod within three seconds to make a rescue shield.",
        "win": "Protect your partner and complete as many team tasks as you can.",
        "safety": "Rescue by touching the pod only. No body contact or blocking paths.",
    },
    {
        "name": "Territory Pulse",
        "tagline": "Capture zones and hold them for your team",
        "players": "4-6",
        "movement": "Medium",
        "image": "05-territory-pulse.png",
        "goal": "Make more floor zones glow in your team's color.",
        "action": "Hold a zone pod until its light changes. Controlled zones earn points over time.",
        "win": "The team that controls zones for the longest time earns the most points.",
        "safety": "Walk between zones and keep every pod where nobody can step on it.",
    },
    {
        "name": "Echo Match",
        "tagline": "See a pattern, then copy it fast",
        "players": "2-6",
        "movement": "Low",
        "image": "06-echo-match.png",
        "goal": "Remember and copy the hub's color or sound pattern.",
        "action": "Watch carefully, then press your own pod in the matching order.",
        "win": "Fast, correct patterns earn points. The most accurate player wins.",
        "safety": "Keep pods stable and use light or vibration if sound is uncomfortable.",
    },
    {
        "name": "Silent Mission",
        "tagline": "Complete a secret task without giving it away",
        "players": "3-6",
        "movement": "Low",
        "image": "07-silent-mission.png",
        "goal": "Finish your hidden mission during a normal conversation or gathering.",
        "action": "Use short or long presses on your own pod to quietly record each mission step.",
        "win": "At the end, the hub reveals which players completed their secret missions.",
        "safety": "Missions must be kind, harmless, and respectful of everyone's space.",
    },
    {
        "name": "Relay Rush",
        "tagline": "Pass the light through the whole team",
        "players": "3-6",
        "movement": "Low-Medium",
        "image": "08-relay-rush.png",
        "goal": "Move a glowing signal through the team as quickly as possible.",
        "action": "When your pod lights, press it to send the signal to the next player.",
        "win": "Finish the route with the lowest time. Mistakes add extra seconds.",
        "safety": "Pods stay attached or on a table - pass the signal, not the hardware.",
    },
    {
        "name": "Crown Circuit",
        "tagline": "Hold the crown and win reaction duels",
        "players": "4-6",
        "movement": "Medium",
        "image": "09-crown-circuit.png",
        "goal": "Keep the digital Crown long enough to collect the most Prestige.",
        "action": "Challenge the Crown pod, step apart, and press your own pod only after the flash.",
        "win": "Hold the Crown, win duels, and reach the Prestige target first.",
        "safety": "Walk in small spaces. Touch only the pod and respect the winner's shield.",
    },
    {
        "name": "Courier Code",
        "tagline": "Hide and pass a secret digital packet",
        "players": "5-6",
        "movement": "Low",
        "image": "10-courier-code.png",
        "goal": "The Courier team secretly passes a packet three times and escapes.",
        "action": "The carrier holds a teammate's pod; that teammate presses their own pod to receive it.",
        "win": "Couriers win by extracting. The Interceptor wins by scanning the real carrier.",
        "safety": "No searching clothes, grabbing hands, or forcing anyone to reveal a clue.",
    },
    {
        "name": "Double Lock",
        "tagline": "Solve clues and press together",
        "players": "3-6",
        "movement": "Low",
        "image": "11-double-lock.png",
        "goal": "Open six digital locks by solving each clue as a team.",
        "action": "Choose the right two or three pods and hold them at the same time.",
        "win": "Open all six locks before the master timer reaches zero.",
        "safety": "Keep pods stable and give every player a clear space to reach.",
    },
    {
        "name": "Market Rush",
        "tagline": "Trade resources and complete contracts",
        "players": "4-6",
        "movement": "Seated",
        "image": "12-market-rush.png",
        "goal": "Trade Spark, Fabric, and Time to finish valuable public contracts.",
        "action": "Build an offer with your pod. The trade happens only when both players confirm.",
        "win": "Completed contracts earn Reputation. The most Reputation wins.",
        "safety": "Everything is pretend - no real money, betting, prizes, or promises.",
    },
    {
        "name": "Power Grid",
        "tagline": "Keep a tiny city powered and cool",
        "players": "3-6",
        "movement": "Low",
        "image": "13-power-grid.png",
        "goal": "Match the city's changing power demand without overheating generators.",
        "action": "Tap for more power, hold to cool, and link two pods for emergency repairs.",
        "win": "Keep Stability above zero until the work shift ends.",
        "safety": "This is a tabletop game. Keep every pod flat, dry, and easy to reach.",
    },
    {
        "name": "Rhythm Rivals",
        "tagline": "Press to the beat and build a team groove",
        "players": "2-6",
        "movement": "Low",
        "image": "14-rhythm-rivals.png",
        "goal": "Hit your notes at the right moment and keep a clean rhythm streak.",
        "action": "Your pod glows before its beat. Press on the bright cue; team notes need everyone together.",
        "win": "The highest Groove score wins, or the whole group reaches its accuracy target.",
        "safety": "Use a comfortable volume and never hit furniture or people for sound.",
    },
    {
        "name": "Switchback",
        "tagline": "Route the signal while the rules change",
        "players": "4-6",
        "movement": "Seated",
        "image": "15-switchback.png",
        "goal": "Keep the live Signal moving to a valid player without making a Fault.",
        "action": "Short, long, and double presses send it different ways - until a modifier changes the rule.",
        "win": "Survive the full round together, or keep the most lives in competitive mode.",
        "safety": "Stay seated. Press your pod; never throw or physically pass the device.",
    },
    {
        "name": "Ghost Frequency",
        "tagline": "Test pairs, connect clues, find the Drifter",
        "players": "4-6",
        "movement": "Low",
        "image": "16-ghost-frequency.png",
        "goal": "Use a few pair tests to discover which player is the hidden Drifter.",
        "action": "Test two pods, compare SAME or DIFFERENT clues, then vote in secret.",
        "win": "Operators win by finding the Drifter. The Drifter wins by staying hidden.",
        "safety": "Accusations are only part of the game - stay kind and include everyone afterward.",
    },
]


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def add_text(slide, text, x, y, w, h, size=20, color=INK, bold=False,
             font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
             margin=0.06, line_spacing=1.0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.clear()
    box.text_frame.word_wrap = True
    box.text_frame.margin_left = Inches(margin)
    box.text_frame.margin_right = Inches(margin)
    box.text_frame.margin_top = Inches(margin)
    box.text_frame.margin_bottom = Inches(margin)
    box.text_frame.vertical_anchor = valign
    p = box.text_frame.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_round_rect(slide, x, y, w, h, fill, radius_shape=MSO_SHAPE.ROUNDED_RECTANGLE,
                   line=None, line_width=1):
    shape = slide.shapes.add_shape(radius_shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def add_picture_cover(slide, path, x, y, w, h, alt_text):
    with Image.open(path) as image:
        image_ratio = image.width / image.height
    box_ratio = w / h
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if image_ratio > box_ratio:
        visible = box_ratio / image_ratio
        pic.crop_left = pic.crop_right = (1 - visible) / 2
    elif image_ratio < box_ratio:
        visible = image_ratio / box_ratio
        pic.crop_top = pic.crop_bottom = (1 - visible) / 2
    pic._element.nvPicPr.cNvPr.set("descr", alt_text)
    return pic


def add_top_bar(slide, section, number=None):
    bar = add_round_rect(slide, 0, 0, SLIDE_W, 0.62, NAVY, MSO_SHAPE.RECTANGLE)
    add_text(slide, "PHYSICAL SOCIAL GAMES", 0.42, 0.14, 4.1, 0.28, 10, WHITE, True)
    label = section if number is None else f"{section}  /  {number:02d}"
    add_text(slide, label.upper(), 9.0, 0.14, 3.9, 0.28, 10, "7FDBD4", True, align=PP_ALIGN.RIGHT)
    return bar


def add_pill(slide, text, x, y, w, fill=MINT, color=TEAL):
    add_round_rect(slide, x, y, w, 0.42, fill, line=color, line_width=0.8)
    add_text(slide, text.upper(), x + 0.08, y + 0.08, w - 0.16, 0.23, 9.5, color, True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)


def add_info_card(slide, label, body, x, y, w, fill, accent):
    add_round_rect(slide, x, y, w, 1.0, fill)
    add_round_rect(slide, x, y, 0.10, 1.0, accent, MSO_SHAPE.RECTANGLE)
    add_text(slide, label.upper(), x + 0.20, y + 0.11, w - 0.32, 0.2, 9, accent, True, margin=0)
    add_text(slide, body, x + 0.20, y + 0.33, w - 0.32, 0.59, 13.2, INK, False, margin=0, line_spacing=0.92)


def add_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(NAVY)
    add_round_rect(slide, 0.55, 0.55, 5.35, 6.4, "163954")
    add_text(slide, "16 GAMES.\nONE MAGIC BOX.", 0.95, 1.0, 4.5, 1.45, 30, WHITE, True,
             line_spacing=0.88)
    add_text(slide, "A visual guide that makes every game easy to understand.",
             0.96, 2.72, 4.45, 0.92, 18, "CFE8E8", False)
    add_round_rect(slide, 0.95, 4.05, 3.7, 0.58, AMBER)
    add_text(slide, "LOOK  •  PRESS  •  PLAY", 1.08, 4.19, 3.45, 0.24, 12, NAVY, True,
             align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, "Child-friendly explanation deck\nCurrent prototype guidance remains supervised 14+.",
             0.96, 5.3, 4.45, 0.75, 11, "AFC3D4")
    add_round_rect(slide, 6.2, 0.55, 6.58, 6.4, WHITE)
    add_picture_cover(slide, HANDBOOK_ASSETS / "kit-concept.png", 6.35, 0.72, 6.28, 5.45,
                      "Concept image of six wearable game pods, one hub, and a carry case")
    add_text(slide, "Six pods + one hub = many kinds of play", 6.65, 6.27, 5.7, 0.36,
             15, NAVY, True, align=PP_ALIGN.CENTER, margin=0)
    return slide


def add_how_it_works_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(PALE)
    add_top_bar(slide, "How it works")
    add_text(slide, "The same buttons become different games", 0.55, 0.88, 8.8, 0.55, 27, NAVY, True)
    add_text(slide, "The pod only tells the hub what you did. The game decides what that action means.",
             0.58, 1.44, 11.9, 0.45, 15, GRAY)
    cards = [
        ("1", "WEAR OR PLACE", "Clip a pod safely, hold it, or place it flat on a table.", TEAL, MINT),
        ("2", "WATCH THE CUE", "The hub and pod show whose turn it is and what is happening.", CORAL, PEACH),
        ("3", "USE ONE ACTION", "Short press, long hold, double press, or press together.", AMBER, CREAM),
    ]
    for i, (num, title, body, accent, fill) in enumerate(cards):
        x = 0.62 + i * 4.18
        add_round_rect(slide, x, 2.18, 3.75, 3.75, WHITE, line="D4DEE7", line_width=1)
        add_round_rect(slide, x + 1.36, 2.48, 1.02, 1.02, accent, MSO_SHAPE.OVAL)
        add_text(slide, num, x + 1.36, 2.66, 1.02, 0.45, 24, WHITE, True,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
        add_text(slide, title, x + 0.35, 3.84, 3.05, 0.36, 14, accent, True,
                 align=PP_ALIGN.CENTER, margin=0)
        add_text(slide, body, x + 0.42, 4.42, 2.9, 0.86, 15, INK,
                 align=PP_ALIGN.CENTER, margin=0, line_spacing=1.02)
    add_round_rect(slide, 0.62, 6.36, 12.05, 0.6, NAVY)
    add_text(slide, "The hub checks every press, keeps the score, and tells every pod what to show.",
             0.95, 6.52, 11.4, 0.24, 13, WHITE, True, align=PP_ALIGN.CENTER, margin=0)


def add_safety_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(PALE)
    add_top_bar(slide, "Safety first")
    add_text(slide, "Easy to understand does not mean ready for age 10", 0.55, 0.86, 11.9, 0.58, 27, NAVY, True)
    add_picture_cover(slide, HANDBOOK_ASSETS / "gameplay-villa.png", 0.58, 1.62, 7.2, 4.63,
                      "Adults playing safely with wearable pods in a clear garden")
    rules = [
        ("CURRENT PLAN", "The prototype is supervised and intended for ages 14+ until safety testing is complete.", CORAL, PEACH),
        ("TOUCH ONLY THE POD", "No pushing, grabbing, tackling, blocking doors, or touching faces and necks.", TEAL, MINT),
        ("CLEAR THE SPACE", "Stay away from roads, stairs, pools, kitchens, vehicles, and breakable objects.", AMBER, CREAM),
        ("STOP RIGHT AWAY", "Stop for pain, heat, damaged hardware, a loose strap, or unsafe behavior.", BLUE, SKY),
    ]
    for i, (title, body, accent, fill) in enumerate(rules):
        y = 1.62 + i * 1.18
        add_round_rect(slide, 8.05, y, 4.65, 0.98, fill)
        add_text(slide, title, 8.26, y + 0.12, 4.15, 0.2, 9.5, accent, True, margin=0)
        add_text(slide, body, 8.26, y + 0.37, 4.15, 0.47, 13.2, INK, margin=0, line_spacing=0.93)
    add_text(slide, "A host checks the area and explains the rules before every round.",
             0.75, 6.48, 11.9, 0.32, 14, NAVY, True, align=PP_ALIGN.CENTER, margin=0)


def add_game_slide(prs, game, index):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(PALE)
    add_top_bar(slide, "Game library", index)

    add_text(slide, game["name"], 0.55, 0.82, 6.1, 0.52, 27, NAVY, True)
    add_text(slide, game["tagline"], 0.58, 1.32, 6.4, 0.34, 13.5, GRAY)
    add_pill(slide, f'{game["players"]} PLAYERS', 8.36, 0.88, 1.48)
    add_pill(slide, game["movement"], 9.98, 0.88, 1.38, CREAM, AMBER)
    add_pill(slide, "ONE HUB", 11.50, 0.88, 1.22, SKY, BLUE)

    add_round_rect(slide, 0.50, 1.82, 7.55, 4.78, WHITE, line="D4DEE7", line_width=1)
    add_picture_cover(slide, ASSET_DIR / game["image"], 0.61, 1.93, 7.33, 4.56,
                      f'Visual explanation of {game["name"]}')

    add_info_card(slide, "Goal", game["goal"], 8.30, 1.82, 4.48, MINT, TEAL)
    add_info_card(slide, "What you do", game["action"], 8.30, 3.02, 4.48, PEACH, CORAL)
    add_info_card(slide, "How to win", game["win"], 8.30, 4.22, 4.48, CREAM, AMBER)

    add_round_rect(slide, 8.30, 5.46, 4.48, 1.14, NAVY)
    add_text(slide, "SAFE PLAY", 8.52, 5.60, 1.05, 0.18, 9, "7FDBD4", True, margin=0)
    add_text(slide, game["safety"], 8.52, 5.87, 3.99, 0.55, 12.5, WHITE, margin=0, line_spacing=0.92)
    add_text(slide, f"{index:02d}", 0.62, 6.82, 0.44, 0.23, 9, GRAY, True, margin=0)
    add_text(slide, "Visual learning guide • Rules simplified for explanation", 8.20, 6.82, 4.55, 0.23,
             8.5, GRAY, align=PP_ALIGN.RIGHT, margin=0)


def add_chooser_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(PALE)
    add_top_bar(slide, "Pick your game")
    add_text(slide, "What kind of game do you feel like playing?", 0.55, 0.88, 11.8, 0.56, 27, NAVY, True)
    groups = [
        ("MOVE & CHASE", "Pulse Hunt\nVault Heist\nGuardian Link\nTerritory Pulse\nCrown Circuit", CORAL, PEACH),
        ("THINK TOGETHER", "Signal Chain\nRelay Rush\nDouble Lock\nPower Grid", TEAL, MINT),
        ("TALK & BLUFF", "Silent Mission\nCourier Code\nMarket Rush\nGhost Frequency", AMBER, CREAM),
        ("FAST HANDS", "Echo Match\nRhythm Rivals\nSwitchback", BLUE, SKY),
    ]
    for i, (title, games, accent, fill) in enumerate(groups):
        x = 0.60 + i * 3.15
        add_round_rect(slide, x, 1.83, 2.83, 4.6, WHITE, line="D4DEE7", line_width=1)
        add_round_rect(slide, x, 1.83, 2.83, 0.72, accent, MSO_SHAPE.RECTANGLE)
        add_text(slide, title, x + 0.15, 2.07, 2.53, 0.24, 12, WHITE, True,
                 align=PP_ALIGN.CENTER, margin=0)
        add_text(slide, games, x + 0.32, 2.92, 2.18, 2.72, 17, NAVY, True,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0, line_spacing=1.15)
        add_round_rect(slide, x + 0.96, 5.83, 0.90, 0.22, fill)
    add_round_rect(slide, 1.47, 6.70, 10.38, 0.45, NAVY)
    add_text(slide, "One box can feel like a chase, puzzle, party, rhythm game, or strategy game.",
             1.70, 6.81, 9.9, 0.20, 12, WHITE, True, align=PP_ALIGN.CENTER, margin=0)


def set_document_properties(prs):
    props = prs.core_properties
    props.title = "Physical Social Games - Visual Guide for Young Learners"
    props.subject = "Child-friendly visual explanation of all sixteen game concepts"
    props.author = "Physical Social Games Project"
    props.keywords = "physical social games, visual guide, game library, wearable pods"
    props.comments = "Explanations are simplified for comprehension. Current prototype guidance remains supervised 14+."


def build():
    for game in GAMES:
        path = ASSET_DIR / game["image"]
        if not path.exists():
            raise FileNotFoundError(path)

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    set_document_properties(prs)

    add_title_slide(prs)
    add_how_it_works_slide(prs)
    add_safety_slide(prs)
    for index, game in enumerate(GAMES, 1):
        add_game_slide(prs, game, index)
    add_chooser_slide(prs)

    prs.save(OUTPUT)
    print(f"Saved {OUTPUT} with {len(prs.slides)} slides")


if __name__ == "__main__":
    build()
